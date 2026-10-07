"""Descarga en Excel con el formato del equipo."""

from io import BytesIO

from fastapi.testclient import TestClient
import pytest
from openpyxl import load_workbook

from app.clientes import CLIENTES
from app.dependencias import obtener_cliente, obtener_fuente
from app.main import app
from app.servicios import excel

INSPECCION = [
    {"ID": "a1", "FECHA": "6/10/2026", "TURNO": "1ro", "HORA": "8:00:00", "NUMERO DE PARTE": "P8MB863242K  BDE",
     "SERIAL": "S1", "FECHA DE PRODUCCION": "1/10/2026", "PIEZAS INSP.": "14", "PIEZAS NOK": "2", "COMENTARIOS": "ok"},
    {"ID": "b2", "FECHA": "7/10/2026", "TURNO": "3ro", "HORA": "2:30:00", "NUMERO DE PARTE": "P8MB863242L ITC",
     "SERIAL": "S2", "FECHA DE PRODUCCION": "", "PIEZAS INSP.": "16", "PIEZAS NOK": "0", "COMENTARIOS": ""},
    {"ID": "c3", "FECHA": "8/10/2026", "TURNO": "1ro", "HORA": "9:00:00", "NUMERO DE PARTE": "P8MB863242K BDE",
     "SERIAL": "S3", "FECHA DE PRODUCCION": "2/10/2026", "PIEZAS INSP.": "10", "PIEZAS NOK": "1", "COMENTARIOS": ""},
]
DEFECTOS = [
    {"ID INSPECCION": "a1", "DEFECTO": "DAÑO", "CANTIDAD": "1"},
    {"ID INSPECCION": "a1", "DEFECTO": "RAYONES", "CANTIDAD": "1"},
    {"ID INSPECCION": "c3", "DEFECTO": "RAYONES", "CANTIDAD": "1"},
]


class Falsa:
    def leer_tabla(self, pestana):
        return DEFECTOS if "Defecto" in pestana else INSPECCION


def descargar(query):
    app.dependency_overrides[obtener_fuente] = lambda: Falsa()
    app.dependency_overrides[obtener_cliente] = lambda: next(iter(CLIENTES.values()))
    try:
        with TestClient(app) as c:
            return c.get("/api/excel?" + query)
    finally:
        app.dependency_overrides.clear()


def hoja_de(res):
    return load_workbook(BytesIO(res.content)).active


def test_columnas_en_el_orden_del_formato():
    res = descargar("desde=2026-10-06&hasta=2026-10-08")
    assert res.status_code == 200
    assert 'filename="inspeccion_hbpo_2026-10-06_a_2026-10-08.xlsx"' in res.headers["content-disposition"]
    encabezado = [c.value for c in hoja_de(res)[1]]
    assert encabezado == [
        "FECHA", "HORA", "TURNO", "NUMERO DE PARTE", "CANTIDAD INSPECCIONADA", "DIA PRODUCCIÓN", "SERIAL",
        "DAÑO", "RAYONES", "CANTIDAD DE PIEZAS INSPECCIONADAS", "SCRAP", "%SCRAP", "TOTAL DE PIEZAS OK",
        "COMENTARIOS",
    ]


def test_valores_y_fila_total():
    h = hoja_de(descargar("desde=2026-10-06&hasta=2026-10-08"))
    fila = [c.value for c in h[2]]
    assert fila[0].date().isoformat() == "2026-10-06"       # fecha real de captura
    assert fila[3] == "8MB863242KBDE"
    assert fila[4] == 14
    assert fila[5].date().isoformat() == "2026-10-01"       # fecha de producción del lote
    assert fila[7:13] == pytest.approx([1, 1, 14, 2, 2 / 14, 12])
    assert h.cell(row=3, column=6).value is None             # sin fecha de producción
    total = [c.value for c in h[5]]
    assert total[0] == "Total"
    assert total[4] == 40
    assert total[7:13] == pytest.approx([1, 2, 40, 3, 3 / 40, 37])
    assert h.auto_filter.ref == "A1:N4"


def test_filtra_por_dia_del_tablero():
    # El registro de las 2:30 del 7 cuenta para el día 6.
    h = hoja_de(descargar("desde=2026-10-06&hasta=2026-10-06"))
    assert [c.value for c in h["G"][1:3]] == ["S1", "S2"]
    assert h.cell(row=4, column=1).value == "Total"


def test_filtros_del_tablero():
    h = hoja_de(descargar("desde=2026-10-06&hasta=2026-10-08&defecto=rayones&parte=kbde"))
    assert [c.value for c in h["G"][1:3]] == ["S1", "S3"]
    h = hoja_de(descargar("desde=2026-10-06&hasta=2026-10-08&turno=3ro"))
    assert h.cell(row=2, column=7).value == "S2"


def test_rango_al_reves_es_error():
    assert descargar("desde=2026-10-08&hasta=2026-10-06").status_code == 400


def test_sin_registros_solo_encabezado_y_total():
    h = load_workbook(BytesIO(excel.armar_libro([], ["DAÑO"]))).active
    assert h.cell(row=2, column=1).value == "Total"
