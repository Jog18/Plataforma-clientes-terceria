"""Reporte de inspección en Excel (.xlsx) con el formato de Jesús (2026-10-07).

Columnas, en orden:
    FECHA, HORA, TURNO, NUMERO DE PARTE, CANTIDAD INSPECCIONADA,
    FECHA DE PRODUCCION, SERIAL, <una columna por defecto del catálogo>,
    CANTIDAD DE PIEZAS INSPECCIONADAS, SCRAP, %SCRAP, TOTAL DE PIEZAS OK,
    COMENTARIOS
y al final una fila "Total".

Después va la hoja "Resumen" (Jesús, 2026-10-08): KPIs, tabla OK/Scrap,
tabla de defectos y dos gráficas (dona y barras). Todo son fórmulas que
apuntan a la fila Total de la hoja de datos, nada queda escrito a mano.

FECHA y HORA son las reales de captura. FECHA DE PRODUCCION es la fecha en que se
produjo el lote inspeccionado (misma columna de la hoja).
Los registros se eligen por el día del tablero (de 6:00 a 5:59).
Reglas HBPO: scrap = NOK, OK = inspeccionadas - NOK.
"""

from datetime import date, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.chart.data_source import StrData, StrRef, StrVal
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Colores del formato (los de la hoja de Excel que usa el equipo).
AZUL = "5B9BD5"
AZUL_OSCURO = "2E75B6"
BLANCO = "FFFFFF"
AMARILLO = "FFFF00"
ROJO = "FF0000"
VERDE = "70AD47"
FRANJA = "DDEBF7"

# (título, color de fondo, color de letra, texto vertical, ancho)
COLS_INICIO = [
    ("FECHA", AZUL, BLANCO, False, 12),
    ("HORA", AZUL, BLANCO, False, 10),
    ("TURNO", AZUL, BLANCO, False, 8),
    ("NUMERO DE PARTE", AZUL, BLANCO, False, 18),
    ("CANTIDAD INSPECCIONADA", AZUL, BLANCO, False, 16),
    ("FECHA DE PRODUCCION", AZUL, BLANCO, False, 15),
    ("SERIAL", AZUL, BLANCO, False, 14),
]
COLS_FINAL = [
    ("CANTIDAD DE PIEZAS INSPECCIONADAS", AMARILLO, "000000", True, 7),
    ("SCRAP", ROJO, BLANCO, True, 7),
    ("%SCRAP", ROJO, BLANCO, True, 8),
    ("TOTAL DE PIEZAS OK", VERDE, BLANCO, True, 7),
    ("COMENTARIOS", AZUL_OSCURO, BLANCO, False, 40),
]

LINEA = Side(style="thin", color="A6A6A6")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)


def filtrar(filas: list[dict], desde: str, hasta: str, turno: str = "",
            parte: str = "", defecto: str = "") -> list[dict]:
    """Mismo filtro que el diálogo del tablero.

    desde/hasta: días de producción 'YYYY-MM-DD' (inclusive).
    parte: texto contenido en el número de parte, sin importar mayúsculas.
    defecto: solo registros que tengan piezas con ese defecto.
    """
    parte = parte.strip().upper()
    defecto = defecto.strip().upper()
    salida = []
    for r in filas:
        if not desde <= r["dia_produccion"] <= hasta:
            continue
        if turno and r["turno"] != turno:
            continue
        if parte and parte not in r["parte"].upper():
            continue
        if defecto and not any(d["defecto"] == defecto and d["cantidad"] > 0 for d in r["defectos"]):
            continue
        salida.append(r)
    return salida


def _fecha(texto: str | None):
    """'2026-10-06' -> date; si la hoja no trae fecha de producción, vacío."""
    return date.fromisoformat(texto) if texto else None


def _hora(texto: str):
    """'08:21:42' -> time(8, 21, 42) para que Excel la trate como hora."""
    try:
        h, m, s = (int(p) for p in texto.split(":"))
        return time(h, m, s)
    except ValueError:
        return texto


def armar_libro(filas: list[dict], catalogo_defectos: list[str], titulo_hoja: str = "Inspeccion") -> bytes:
    """Regresa el archivo .xlsx (en bytes) con el formato del reporte."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo_hoja[:31]

    columnas = (COLS_INICIO
                + [(d, BLANCO, ROJO, True, 7) for d in catalogo_defectos]
                + COLS_FINAL)

    # Encabezado.
    for c, (nombre, fondo, letra, vertical, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=c, value=nombre)
        celda.fill = PatternFill("solid", fgColor=fondo)
        celda.font = Font(bold=True, color=letra, size=10)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True,
                                    text_rotation=90 if vertical else 0)
        celda.border = BORDE
        hoja.column_dimensions[get_column_letter(c)].width = ancho
    hoja.row_dimensions[1].height = 120

    n_def = len(catalogo_defectos)
    col_def = len(COLS_INICIO) + 1          # primera columna de defecto
    col_insp2 = col_def + n_def             # CANTIDAD DE PIEZAS INSPECCIONADAS
    col_pct = col_insp2 + 2                 # %SCRAP
    idx_def = {d: i for i, d in enumerate(catalogo_defectos)}

    tot_def = [0] * n_def
    tot_insp = tot_nok = tot_ok = 0
    for i, r in enumerate(filas):
        por_def = [0] * n_def
        for d in r["defectos"]:
            if d["defecto"] in idx_def:
                por_def[idx_def[d["defecto"]]] += d["cantidad"]
        insp, nok, ok = r["insp"], r["nok"], r["ok"]
        valores = [
            date.fromisoformat(r["fecha"]), _hora(r["hora"]), r["turno"], r["parte"], insp,
            _fecha(r["fecha_produccion"]), r["serial"],
            *por_def,
            insp, r["scrap"], (nok / insp) if insp else 0, ok, r["comentarios"],
        ]
        fila = i + 2
        franja = PatternFill("solid", fgColor=FRANJA) if i % 2 else None
        for c, v in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=c, value=v)
            celda.border = BORDE
            if franja:
                celda.fill = franja
            if c != len(columnas):
                celda.alignment = Alignment(horizontal="center")
        hoja.cell(row=fila, column=1).number_format = "d/m/yyyy"
        hoja.cell(row=fila, column=2).number_format = "hh:mm:ss"
        hoja.cell(row=fila, column=6).number_format = "d/m/yyyy"
        hoja.cell(row=fila, column=col_pct).number_format = "0.00%"

        tot_def = [a + b for a, b in zip(tot_def, por_def)]
        tot_insp += insp
        tot_nok += nok
        tot_ok += ok

    # Fila "Total".
    fila = len(filas) + 2
    totales = {1: "Total", 5: tot_insp}
    totales.update({col_def + j: v for j, v in enumerate(tot_def)})
    totales.update({col_insp2: tot_insp, col_insp2 + 1: tot_nok,
                    col_pct: (tot_nok / tot_insp) if tot_insp else 0, col_pct + 1: tot_ok})
    for c in range(1, len(columnas) + 1):
        celda = hoja.cell(row=fila, column=c, value=totales.get(c))
        celda.font = Font(bold=True)
        celda.border = BORDE
        celda.alignment = Alignment(horizontal="center")
    hoja.cell(row=fila, column=col_pct).number_format = "0.00%"

    # Flechas de filtro sobre los registros (sin la fila Total) y encabezado fijo.
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(fila - 1, 1)}"
    hoja.freeze_panes = "A2"
    # Al imprimir: horizontal y todas las columnas en una hoja de ancho.
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = "1:1"

    _hoja_resumen(libro, hoja, fila, col_def, catalogo_defectos, col_insp2, (tot_insp, tot_ok, tot_nok))

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---- hoja "Resumen" -------------------------------------------------------------

AZUL_RESUMEN = "1F3864"
GRIS_PERIODO = "595959"
GRIS_TITULO = "808080"
VERDE_KPI = "008000"
VERDE_OK = "2E7D32"
ROJO_SCRAP = "C62828"
LINEA_RESUMEN = Side(style="thin", color="000000")

# Anchos pedidos en puntos; Excel mide en caracteres (~7 px, 1 pt = 4/3 px).
ANCHOS_PUNTOS = {"A": 15, "B": 150, "C": 70, "D": 15, "E": 95, "F": 60}


def nombre_defecto_resumen(defecto: str) -> str:
    """'DAÑO' -> 'Daño', 'LÁSER NOK' -> 'Defecto láser'."""
    if "LÁSER" in defecto.upper():
        return "Defecto láser"
    return defecto.capitalize()


def orden_defectos_resumen(catalogo: list[str]) -> list[str]:
    """Daño, Rayones y láser primero (orden de Jesús); los demás después."""
    def clave(d: str):
        d = d.upper()
        if d.startswith("DAÑO"):
            return 0
        if d.startswith("RAYON"):
            return 1
        if "LÁSER" in d:
            return 2
        return 3
    return sorted(catalogo, key=clave)


def _texto_grafica(tamano: int, color: str, negrita: bool = True) -> RichText:
    props = CharacterProperties(sz=tamano * 100, b=negrita, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])


def _encabezado_azul(hoja, rango: str, textos: tuple[str, str]) -> None:
    for celda, texto in zip(hoja[rango][0], textos):
        celda.value = texto
        celda.fill = PatternFill("solid", fgColor=AZUL_RESUMEN)
        celda.font = Font(bold=True, color=BLANCO)
        celda.border = Border(bottom=LINEA_RESUMEN)


def _hoja_resumen(libro, datos, fila_total: int, col_def: int,
                  catalogo: list[str], col_insp2: int, totales: tuple[int, int, int]) -> None:
    h = libro.create_sheet("Resumen")       # queda justo después de la de datos
    ref = "'" + datos.title.replace("'", "''") + "'!"
    letra = get_column_letter
    ultima = max(fila_total - 1, 2)
    fechas = f"{ref}A2:A{ultima}"

    for col, puntos in ANCHOS_PUNTOS.items():
        h.column_dimensions[col].width = round(puntos * 4 / 3 / 7, 2)

    # 1. Encabezado.
    h["B2"] = f"Resumen de Inspección {datos.title} – Total del periodo"
    h["B2"].font = Font(name="Calibri", size=16, bold=True, color=AZUL_RESUMEN)
    dia = lambda f: (f'TEXT(DAY({f}({fechas})),"00")&"/"&TEXT(MONTH({f}({fechas})),"00")'
                     f'&"/"&YEAR({f}({fechas}))')
    h["B3"] = f'="Periodo: "&{dia("MIN")}&" al "&{dia("MAX")}'
    h["B3"].font = Font(italic=True, size=11, color=GRIS_PERIODO)

    # 2. KPIs.
    _encabezado_azul(h, "B5:C5", ("Indicador", "Valor"))
    kpis = [
        ("Piezas inspeccionadas", f"={ref}{letra(col_insp2)}{fila_total}"),
        ("Piezas OK", f"={ref}{letra(col_insp2 + 3)}{fila_total}"),
        ("Piezas scrap", f"={ref}{letra(col_insp2 + 1)}{fila_total}"),
        ("% Scrap", "=IF(C6=0,0,C8/C6)"),
    ]
    for i, (nombre, formula) in enumerate(kpis, start=6):
        h.cell(row=i, column=2, value=nombre)
        c = h.cell(row=i, column=3, value=formula)
        if i < 9:
            c.font = Font(color=VERDE_KPI)
            c.number_format = "#,##0"
    h["C9"].number_format = "0.00%"
    h["B9"].font = h["C9"].font = Font(bold=True)
    for fila_kpi in h["B6:C9"]:
        for c in fila_kpi:
            c.border = Border(top=LINEA_RESUMEN, bottom=LINEA_RESUMEN)

    # 3. Resultado (fuente de la dona).
    _encabezado_azul(h, "E5:F5", ("Resultado", "Piezas"))
    for i, (nombre, formula) in enumerate([("OK", "=C7"), ("Scrap", "=C8")], start=6):
        h.cell(row=i, column=5, value=nombre)
        h.cell(row=i, column=6, value=formula).number_format = "#,##0"

    # 4. Defectos (fuente de las barras), uno por defecto del catálogo.
    _encabezado_azul(h, "E10:F10", ("Tipo de defecto", "Piezas"))
    orden = orden_defectos_resumen(catalogo)
    for i, d in enumerate(orden, start=11):
        col = col_def + catalogo.index(d)
        h.cell(row=i, column=5, value=nombre_defecto_resumen(d))
        h.cell(row=i, column=6, value=f"={ref}{letra(col)}{fila_total}").number_format = "#,##0"
    fin_def = 10 + max(len(orden), 1)

    # 5. Título dinámico de la dona.
    h["E3"] = ('="OK: "&TEXT(C7,"#,##0")&" ("&TEXT(IF(C6=0,0,C7/C6),"0.0%")&") | Scrap: "'
               '&TEXT(C8,"#,##0")&" ("&TEXT(C9,"0.00%")&")"')
    h["E3"].font = Font(size=9, color=GRIS_TITULO)

    # 6. Dona.
    dona = DoughnutChart(holeSize=60)
    dona.add_data(Reference(h, min_col=6, min_row=5, max_row=7), titles_from_data=True)
    dona.set_categories(Reference(h, min_col=5, min_row=6, max_row=7))
    serie = dona.series[0]
    for idx, color in enumerate((VERDE_OK, ROJO_SCRAP)):
        serie.dPt.append(DataPoint(idx=idx, spPr=GraphicalProperties(solidFill=color)))
    serie.dLbls = DataLabelList(showCatName=True, showVal=True, showPercent=False,
                                showSerName=False, showLeaderLines=False, showLegendKey=False,
                                separator=": ", numFmt="#,##0",
                                txPr=_texto_grafica(10, "000000"))
    dona.legend.position = "b"
    # El título se vincula a E3; el texto guardado es solo la vista previa
    # hasta que Excel recalcula.
    insp, ok, nok = totales
    pct = lambda a, d: f"{(a / insp if insp else 0) * 100:.{d}f}%"
    previo = f"OK: {ok:,} ({pct(ok, 1)}) | Scrap: {nok:,} ({pct(nok, 2)})"
    cache = StrData(ptCount=1, pt=[StrVal(idx=0, v=previo)])
    dona.title = Title(tx=Text(strRef=StrRef("'Resumen'!$E$3", strCache=cache)), overlay=False,
                       txPr=_texto_grafica(13, AZUL_RESUMEN))
    dona.width, dona.height = 16, 11.5
    h.add_chart(dona, "H2")

    # 7. Barras horizontales.
    barras = BarChart(barDir="bar", gapWidth=60)
    barras.add_data(Reference(h, min_col=6, min_row=10, max_row=fin_def), titles_from_data=True)
    barras.set_categories(Reference(h, min_col=5, min_row=11, max_row=fin_def))
    barras.series[0].graphicalProperties = GraphicalProperties(solidFill=ROJO_SCRAP)
    barras.series[0].dLbls = DataLabelList(showVal=True, showCatName=False, showSerName=False,
                                           showLegendKey=False, showPercent=False,
                                           txPr=_texto_grafica(9, "000000"))
    barras.legend = None
    barras.title = "Scrap por tipo de defecto (piezas)"
    barras.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1300, b=True))
    barras.title.overlay = False                     # título arriba, sin tapar el eje
    barras.x_axis.scaling.orientation = "maxMin"     # Daño arriba
    barras.y_axis.majorGridlines = None
    barras.y_axis.scaling.min = 0
    barras.x_axis.delete = False
    barras.y_axis.delete = False
    barras.width, barras.height = 15, 7.5
    h.add_chart(barras, "B15")
